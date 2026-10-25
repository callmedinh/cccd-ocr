"""Ghi kết quả ra Excel (.xlsx) và CSV (mở trực tiếp bằng Google Sheets)."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Sequence
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import CardRecord

HEADERS: tuple[str, ...] = (
    "STT", "Số CCCD", "Họ và tên", "Ngày sinh", "Giới tính", "Quốc tịch",
    "Quê quán", "Nơi thường trú", "Ngày cấp", "Có giá trị đến",
    "Đặc điểm nhận dạng", "Nơi cấp", "Độ tin cậy", "Cảnh báo", "Tệp nguồn",
)

COLUMN_WIDTHS: dict[str, int] = {
    "STT": 6, "Số CCCD": 16, "Họ và tên": 26, "Ngày sinh": 13, "Giới tính": 10,
    "Quốc tịch": 12, "Quê quán": 38, "Nơi thường trú": 48, "Ngày cấp": 13,
    "Có giá trị đến": 14, "Đặc điểm nhận dạng": 30, "Nơi cấp": 34,
    "Độ tin cậy": 11, "Cảnh báo": 40, "Tệp nguồn": 30,
}

# Cột dạng ngày: giữ định dạng text để Excel không tự đổi dd/mm <-> mm/dd
TEXT_COLUMNS = {"Số CCCD", "Ngày sinh", "Ngày cấp", "Có giá trị đến"}

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def build_workbook(records: Sequence[CardRecord], *, sheet_title: str = "CCCD",
                   start_index: int = 0):
    """Dựng ``Workbook`` đã định dạng sẵn từ danh sách bản ghi."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append(list(HEADERS))
    _style_header(sheet)
    append_rows(sheet, records, start_index=start_index)
    _finish_sheet(sheet)
    return workbook


def append_rows(sheet: Worksheet, records: Sequence[CardRecord], *,
                start_index: int = 0) -> None:
    """Thêm các bản ghi vào sheet đã có tiêu đề và tô định dạng."""
    first_row = sheet.max_row + 1
    for offset, record in enumerate(records):
        values = record.to_row(start_index + offset + 1)
        sheet.append([values[header] for header in HEADERS])
    for row_index in range(first_row, sheet.max_row + 1):
        _style_row(sheet, row_index)


def write_excel(records: Sequence[CardRecord], output_path: str | Path,
                *, sheet_title: str = "CCCD", append: bool = False) -> Path:
    """Ghi danh sách bản ghi ra file Excel có định dạng sẵn để đọc/in."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if append and path.exists():
        workbook = load_workbook(path)
        sheet = workbook[sheet_title] if sheet_title in workbook.sheetnames else workbook.active
        append_rows(sheet, records, start_index=sheet.max_row)
        _finish_sheet(sheet)
    else:
        workbook = build_workbook(records, sheet_title=sheet_title)

    workbook.save(path)
    return path


def workbook_bytes(records: Sequence[CardRecord], *, sheet_title: str = "CCCD") -> bytes:
    """Như ``write_excel`` nhưng trả về bytes (dùng cho tải xuống qua HTTP)."""
    workbook = build_workbook(records, sheet_title=sheet_title)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def csv_bytes(records: Iterable[CardRecord]) -> bytes:
    """Nội dung CSV UTF-8 (có BOM) dưới dạng bytes."""
    text = io.StringIO()
    writer = csv.DictWriter(text, fieldnames=list(HEADERS))
    writer.writeheader()
    for index, record in enumerate(records, start=1):
        writer.writerow(record.to_row(index))
    return text.getvalue().encode("utf-8-sig")


def _style_header(sheet: Worksheet) -> None:
    for column_index, _header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=column_index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    sheet.row_dimensions[1].height = 26


def _style_row(sheet: Worksheet, row_index: int) -> None:
    has_warning = bool(sheet.cell(row=row_index, column=HEADERS.index("Cảnh báo") + 1).value)
    for column_index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=row_index, column=column_index)
        cell.border = _BORDER
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=header in ("Quê quán", "Nơi thường trú", "Đặc điểm nhận dạng",
                                 "Nơi cấp", "Cảnh báo"),
            horizontal="center" if header in ("STT", "Giới tính", "Độ tin cậy", "Số CCCD")
            else "left",
        )
        if header in TEXT_COLUMNS:
            cell.number_format = "@"   # ép text: tránh Excel cắt số 0 đầu & đảo ngày
        if has_warning and header == "Cảnh báo":
            cell.fill = _WARN_FILL


def _finish_sheet(sheet: Worksheet) -> None:
    for column_index, header in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = COLUMN_WIDTHS.get(
            header, 18
        )
    sheet.freeze_panes = "B2"          # giữ dòng tiêu đề + cột STT khi cuộn
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(sheet.max_row, 1)}"


def write_csv(records: Iterable[CardRecord], output_path: str | Path) -> Path:
    """Ghi CSV UTF-8 (có BOM) — mở bằng Google Sheets/Excel đều đúng tiếng Việt."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(HEADERS))
        writer.writeheader()
        for index, record in enumerate(records, start=1):
            writer.writerow(record.to_row(index))
    return path
